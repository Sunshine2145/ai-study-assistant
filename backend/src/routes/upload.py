# AI Study Assistant - Upload Routes

import os
import json
import uuid
from fastapi import APIRouter, UploadFile, File, HTTPException
from pydantic import BaseModel
from typing import Optional, List

from src.database.db import db
from src.config.settings import settings

router = APIRouter(prefix="/api/upload", tags=["上传"])


class QuestionItem(BaseModel):
    type: str
    content: str
    options: Optional[dict] = {}
    answer: str
    analysis: Optional[str] = ""
    difficulty: Optional[int] = 1


class QuestionImportRequest(BaseModel):
    questions: List[QuestionItem]
    knowledge_point: Optional[str] = None
    source: Optional[str] = "manual"


@router.post("/questions")
async def upload_questions(
    file: UploadFile = File(...),
    knowledge_point: Optional[str] = None,
    use_ai_parse: bool = False,
    ai_provider: str = "deepseek"
):
    """上传题库文件（支持TXT、JSON、MD、PDF、DOCX、XLSX格式）"""
    if not file.filename:
        raise HTTPException(status_code=400, detail="文件名为空")

    # 验证AI提供商
    valid_providers = ["minimax", "deepseek"]
    if ai_provider not in valid_providers:
        ai_provider = "deepseek"

    # 读取文件内容
    content = await file.read()

    questions = []
    filename = file.filename.lower()

    try:
        if filename.endswith('.json'):
            # JSON格式直接解析
            questions = json.loads(content.decode('utf-8'))
        elif filename.endswith('.txt'):
            # TXT格式解析
            text = content.decode('utf-8')
            if use_ai_parse:
                questions = await _parse_pdf_with_ai(text, ai_provider)
            else:
                # 使用正则解析TXT文本
                questions = _parse_txt_format(text)
        elif filename.endswith('.md'):
            # Markdown格式解析
            text = content.decode('utf-8')
            if use_ai_parse:
                questions = await _parse_pdf_with_ai(text, ai_provider)
            else:
                # 使用正则解析MD文本
                questions = _parse_markdown_format(text)
        elif filename.endswith('.pdf'):
            # PDF格式解析
            text = await _parse_pdf_content(content)
            if use_ai_parse:
                questions = await _parse_pdf_with_ai(text, ai_provider)
            else:
                # 使用正则解析PDF文本
                questions = _parse_pdf_text(text)
        elif filename.endswith('.docx'):
            # Word文档解析
            text = _parse_docx_content(content)
            if use_ai_parse:
                questions = await _parse_pdf_with_ai(text, ai_provider)
            else:
                questions = _parse_txt_format(text)
        elif filename.endswith(('.xlsx', '.xls')):
            # Excel文件解析
            text = _parse_excel_content(content)
            if use_ai_parse:
                questions = await _parse_pdf_with_ai(text, ai_provider)
            else:
                questions = _parse_txt_format(text)
        else:
            raise HTTPException(status_code=400, detail="不支持的文件格式，请上传 .txt、.json、.md、.pdf、.docx 或 .xlsx 文件")

        if not questions:
            return {"success": True, "imported": 0, "message": "未找到有效题目"}

        # 导入到数据库
        imported_count = _import_questions(questions, knowledge_point, file.filename)

        return {
            "success": True,
            "imported": imported_count,
            "total_found": len(questions),
            "message": f"成功导入 {imported_count} 道题目"
        }

    except json.JSONDecodeError as e:
        raise HTTPException(status_code=400, detail=f"JSON格式错误：{str(e)}")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"处理失败：{str(e)}")


@router.post("/questions/direct")
async def import_questions_direct(request: QuestionImportRequest):
    """直接导入题目列表"""
    imported = _import_questions(
        [q.dict() for q in request.questions],
        request.knowledge_point,
        request.source
    )
    return {
        "success": True,
        "imported": imported,
        "message": f"成功导入 {imported} 道题目"
    }


def _parse_txt_format(text: str) -> list:
    """解析TXT格式题目"""
    questions = []
    # 按空行分割题目
    blocks = text.split('\n\n')

    for block in blocks:
        block = block.strip()
        if not block:
            continue

        lines = block.split('\n')
        if len(lines) < 2:
            continue

        # 第一行是题目内容
        content = lines[0].strip()

        # 尝试识别类型和答案
        q_type = 'single'
        answer = ''
        options = {}

        for line in lines[1:]:
            line = line.strip()
            if not line:
                continue

            # 判断题
            if line.startswith('对') or line.startswith('错') or line.startswith('T') or line.startswith('F'):
                q_type = 'judge'
                answer = 'true' if line[0] in '对T' else 'false'
            # 选择题选项
            elif line and line[0] in 'ABCD' and ':：' in line:
                key = line[0]
                value = line[1:].strip(':：').strip()
                options[key] = value
            # 答案行
            elif line.startswith('答案') or line.startswith('答:'):
                answer = line.split('答')[1].strip(':：').strip()
                if answer in options:
                    pass  # 答案可能是选项字母
                elif answer in ['对', '错', 'T', 'F']:
                    q_type = 'judge'
                    answer = 'true' if answer in '对T' else 'false'

        if content and answer:
            questions.append({
                'type': q_type,
                'content': content,
                'options': options if options else {},
                'answer': answer,
                'analysis': ''
            })

    return questions


async def _parse_pdf_content(content: bytes) -> str:
    """从PDF中提取文本内容"""
    import fitz  # PyMuPDF
    import io

    text_parts = []
    try:
        # 打开PDF文档
        pdf_doc = fitz.open(stream=content, filetype="pdf")

        for page_num in range(len(pdf_doc)):
            page = pdf_doc[page_num]
            text = page.get_text()
            if text.strip():
                text_parts.append(f"=== 第{page_num + 1}页 ===\n{text}")

        pdf_doc.close()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"PDF解析失败：{str(e)}")

    return "\n\n".join(text_parts)


async def _parse_pdf_with_ai(text: str, ai_provider: str = "deepseek") -> list:
    """使用AI解析PDF中的题目"""
    from src.modules.ai.pdf_parser import PDFParser

    if ai_provider == "minimax":
        parser = PDFParser(
            api_key=settings.ai.api_key,
            model=settings.ai.model,
            base_url=settings.ai.base_url
        )
    else:
        # Default to DeepSeek for better Chinese support
        parser = PDFParser(
            api_key=settings.deepseek.api_key,
            model=settings.deepseek.model,
            base_url=settings.deepseek.base_url
        )
    questions = await parser.parse(text)
    return questions


def _parse_docx_content(content: bytes) -> str:
    """从DOCX中提取文本内容"""
    from docx import Document
    import io

    text_parts = []
    try:
        doc = Document(io.BytesIO(content))
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"DOCX解析失败：{str(e)}")

    return "\n\n".join(text_parts)


def _parse_excel_content(content: bytes) -> str:
    """从Excel中提取文本内容"""
    from openpyxl import load_workbook
    import io

    text_parts = []
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text_parts.append(f"=== {sheet_name} ===")
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    text_parts.append(row_text)
        wb.close()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Excel解析失败：{str(e)}")

    return "\n\n".join(text_parts)


def _parse_pdf_text(text: str) -> list:
    """使用正则表达式解析PDF提取的文本"""
    questions = []
    import re

    # 清理文本
    text = re.sub(r'\s+', ' ', text)
    text = re.sub(r'=== 第\d+页 ===', '\n', text)

    # 匹配选择题：题目内容和选项在一起
    # 模式：题目内容\nA. 选项\nB. 选项\nC. 选项\nD. 选项\n答案: X
    single_pattern = r'([^\n答案]+?)\s*\n\s*([A-D])[.、:：]\s*([^\n]+?)\s*\n\s*([A-D])[.、:：]\s*([^\n]+?)\s*\n\s*([A-D])[.、:：]\s*([^\n]+?)\s*\n\s*([A-D])[.、:：]\s*([^\n]+?)\s*\n\s*答案[：:]\s*([A-D])'

    for match in re.finditer(single_pattern, text, re.MULTILINE):
        content = match.group(1).strip()
        options = {
            match.group(2): match.group(3).strip(),
            match.group(4): match.group(5).strip(),
            match.group(6): match.group(7).strip(),
            match.group(8): match.group(9).strip()
        }
        answer = match.group(10)
        questions.append({
            'type': 'single',
            'content': content,
            'options': options,
            'answer': answer,
            'analysis': ''
        })

    # 匹配判断题
    judge_pattern = r'([^\n答案]+?)\s*\n\s*答案[：:]\s*(对|错|√|×|true|false|T|F)'
    for match in re.finditer(judge_pattern, text, re.MULTILINE):
        content = match.group(1).strip()
        answer_raw = match.group(2)
        answer_map = {'对': 'true', '错': 'false', '√': 'true', '×': 'false', 'true': 'true', 'false': 'false', 'T': 'true', 'F': 'false'}
        answer = answer_map.get(answer_raw, 'true')
        questions.append({
            'type': 'judge',
            'content': content,
            'options': {},
            'answer': answer,
            'analysis': ''
        })

    return questions


def _parse_markdown_format(text: str) -> list:
    """解析Markdown格式题目"""
    questions = []

    # 匹配题目块
    import re
    pattern = r'(?:^|\n)(###\s*)?(.+?)\n((?:[-ABCD].+\n)+)(?:答案[：:]\s*(\S+))?'

    matches = re.finditer(pattern, text, re.MULTILINE)

    for match in matches:
        content = match.group(2).strip()
        options_text = match.group(3) or ''
        answer = match.group(4) or ''

        options = {}
        for opt_match in re.finditer(r'^([ABCD])[.、:：]\s*(.+)$', options_text, re.MULTILINE):
            options[opt_match.group(1)] = opt_match.group(2).strip()

        # 判断题型
        q_type = 'single'
        if answer in ['对', '错', 'true', 'false', 'T', 'F']:
            q_type = 'judge'
            answer = 'true' if answer in '对Ttrue' else 'false'
        elif len(options) > 2 and answer not in options:
            # 可能多选题
            q_type = 'multi'

        questions.append({
            'type': q_type,
            'content': content,
            'options': options,
            'answer': answer if answer in options else list(options.keys())[0] if options else 'A',
            'analysis': ''
        })

    return questions


def _import_questions(questions: list, knowledge_point: Optional[str], source: str) -> int:
    """导入题目到数据库"""
    imported = 0

    for q in questions:
        try:
            # 序列化options为JSON字符串
            options_json = json.dumps(q.get('options', {}), ensure_ascii=False)

            db.execute(
                """INSERT INTO questions (type, content, options, answer, analysis, knowledge_point, source, source_file, status)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 'approved')""",
                (
                    q.get('type', 'single'),
                    q.get('content', ''),
                    options_json,
                    str(q.get('answer', '')),
                    q.get('analysis', ''),
                    knowledge_point or '',
                    source,
                    ''
                )
            )
            imported += 1
        except Exception as e:
            # 跳过重复或无效题目
            continue

    return imported


@router.get("/templates")
async def get_import_template():
    """获取导入模板"""
    template = {
        "format": "JSON数组",
        "example": [
            {
                "type": "single",
                "content": "下列关于Cache的说法，错误的是？",
                "options": {"A": "位于CPU和内存之间", "B": "容量比内存大", "C": "访问速度比内存快", "D": "由高速SRAM组成"},
                "answer": "B",
                "analysis": "Cache的容量通常比内存小很多"
            },
            {
                "type": "judge",
                "content": "Cache的命中率越高越好。",
                "options": {},
                "answer": "false",
                "analysis": "命中率太高可能增加复杂度"
            }
        ],
        "fields": {
            "type": "题型：single(单选), multi(多选), judge(判断)",
            "content": "题目内容",
            "options": "选项，JSON对象，key为选项字母",
            "answer": "答案，单选/判断为选项字母或true/false，多选为字母数组",
            "analysis": "解析（可选）"
        },
        "supported_formats": [".json", ".txt", ".md", ".pdf", ".docx", ".xlsx"]
    }
    return template
